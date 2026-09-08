import io
from typing import Tuple
import docx

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc_io = io.BytesIO(file_bytes)
    doc = docx.Document(doc_io)

    text_parts = []

    # Extract paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            text_parts.append(p.text.strip())

    # Extract tables if any
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_parts.append(' | '.join(row_text))

    full_text = '\n\n'.join(text_parts).strip()
    return full_text
